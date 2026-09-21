"""Exports lead records to a formatted .xlsx file."""

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

import config

logger = config.get_logger(__name__)

COLUMNS = [
    "Business Name",
    "Address",
    "Phone Number",
    "Email ID",
    "Website",
    "Instagram",
    "Facebook",
]
HYPERLINK_COLUMNS = ("Website", "Instagram", "Facebook")


def export_to_excel(records: list[dict], filename: str) -> str:
    """Writes records (with keys: name, address, phone, email, website,
    instagram, facebook) to an .xlsx file inside config.OUTPUT_DIR, with
    headers, column widths, and clickable hyperlinks. Returns the full
    output path.
    """
    rows = [
        {
            "Business Name": r.get("name") or "",
            "Address": r.get("address") or "",
            "Phone Number": r.get("phone") or "",
            "Email ID": r.get("email") or "",
            "Website": r.get("website") or "",
            "Instagram": r.get("instagram") or "",
            "Facebook": r.get("facebook") or "",
        }
        for r in records
    ]
    df = pd.DataFrame(rows, columns=COLUMNS)

    output_path = config.OUTPUT_DIR / filename

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Leads")
        worksheet = writer.sheets["Leads"]

        for col_idx, col_name in enumerate(COLUMNS, start=1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.font = Font(bold=True)
            max_len = max([len(col_name)] + [len(str(v)) for v in df[col_name]])
            worksheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

        for col_name in HYPERLINK_COLUMNS:
            col_idx = COLUMNS.index(col_name) + 1
            for row_idx, value in enumerate(df[col_name], start=2):
                if value:
                    worksheet.cell(row=row_idx, column=col_idx).hyperlink = value

    logger.info("Exported %d records to %s", len(records), output_path)
    return str(output_path)
